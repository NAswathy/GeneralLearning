from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from .models import Employee


def _headers(sheet) -> dict[str, int]:
    return {
        str(cell.value).strip().casefold(): cell.column
        for cell in sheet[1]
        if cell.value is not None
    }


def _column(headers: dict[str, int], label: str, workbook_path: Path) -> int:
    try:
        return headers[label.strip().casefold()]
    except KeyError as exc:
        raise ValueError(f"Column '{label}' not found in {workbook_path}") from exc


def _as_date(value, workbook_path: Path, row_number: int, label: str) -> date | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        for date_format in ("%Y-%m-%d", "%m/%d/%Y", "%d/%m/%Y"):
            try:
                return datetime.strptime(value.strip(), date_format).date()
            except ValueError:
                continue
    raise ValueError(f"Invalid {label} at row {row_number} in {workbook_path}: {value!r}")


def load_employees(
    workbook_path: Path,
    sheet_name: str,
    name_column: str,
    birthday_column: str,
    hire_date_column: str,
) -> list[Employee]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        sheet = workbook[sheet_name]
        headers = _headers(sheet)
        name_index = _column(headers, name_column, workbook_path)
        birthday_index = _column(headers, birthday_column, workbook_path)
        hire_date_index = _column(headers, hire_date_column, workbook_path)
        employees = []
        for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            name_value = row[name_index - 1] if len(row) >= name_index else None
            if name_value is None or not str(name_value).strip():
                continue
            birthday_value = row[birthday_index - 1] if len(row) >= birthday_index else None
            hire_date_value = row[hire_date_index - 1] if len(row) >= hire_date_index else None
            employees.append(
                Employee(
                    name=str(name_value).strip(),
                    birthday=_as_date(birthday_value, workbook_path, row_number, birthday_column),
                    hire_date=_as_date(hire_date_value, workbook_path, row_number, hire_date_column),
                )
            )
        return employees
    finally:
        workbook.close()


def load_recipients(workbook_path: Path, sheet_name: str, email_column: str) -> list[str]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        sheet = workbook[sheet_name]
        headers = _headers(sheet)
        email_index = _column(headers, email_column, workbook_path)
        recipients = []
        for row in sheet.iter_rows(min_row=2, values_only=True):
            value = row[email_index - 1] if len(row) >= email_index else None
            if value is not None and str(value).strip():
                recipients.append(str(value).strip())
        return list(dict.fromkeys(recipients))
    finally:
        workbook.close()